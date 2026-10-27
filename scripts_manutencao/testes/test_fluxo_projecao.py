"""Fluxo projetado: apenas fixtures locais; nunca usa banco real."""
from datetime import date, timedelta
from decimal import Decimal
from io import BytesIO
import openpyxl
from flask import template_rendered
from app.extensoes import db
from app.financeiro.financeiro_model import ContaBancaria, LancamentoFinanceiro, ExtratoBancario, ConciliacaoBancariaItem
from app.financeiro.fluxo_projecao_service import projetar_fluxo
from scripts_manutencao.testes.test_pr46_os_recebimento import banco_pr46, caso_jumirim_ficticio


def lanc(ids, descricao, tipo, valor, vencimento, **kwargs):
    defaults = dict(status='pendente', origem='MANUAL', ativo=True,
                    data_lancamento=date(2026,9,1), conta_bancaria_id=ids['conta'])
    defaults.update(kwargs)
    obj = LancamentoFinanceiro(descricao=descricao, tipo=tipo, valor=Decimal(valor),
                               data_vencimento=vencimento, **defaults)
    db.session.add(obj)
    db.session.flush()
    return obj


def test_projecao_diferencia_pago_pendente_conciliado_e_atrasado(caso_jumirim_ficticio):
    app, ids = caso_jumirim_ficticio
    with app.app_context():
        inicio = date(2026,9,10)
        vencimento = date(2026,9,11)
        lanc(ids, 'Receber', 'conta_receber', '1000', vencimento)
        lanc(ids, 'Pagar', 'conta_pagar', '300', vencimento)
        lanc(ids, 'Pago antecipado', 'conta_pagar', '500', vencimento, status='pago', data_pagamento=inicio)
        lanc(ids, 'Cancelado', 'despesa', '100', vencimento, status='cancelado')
        lanc(ids, 'Inativo', 'receita', '800', vencimento, ativo=False)
        lanc(ids, 'Em atraso', 'despesa', '90', inicio - timedelta(days=1))
        parcial = lanc(ids, 'Parcial', 'conta_receber', '1000', vencimento)
        ext = db.session.get(ExtratoBancario, ids['extrato'])
        db.session.add(ConciliacaoBancariaItem(extrato_id=ext.id, lancamento_id=parcial.id,
                                               valor_conciliado=Decimal('250'), ativo=True))
        # Mesmo par em modo legado não deve ser subtraído duas vezes.
        ext.lancamento_id = parcial.id
        ext.conciliado = True
        db.session.commit()
        saldo = db.session.get(ContaBancaria, ids['conta']).saldo_atual
        resultado = projetar_fluxo(inicio, inicio + timedelta(days=30), ids['conta'])
        assert resultado['saldo_inicial'] == saldo
        assert resultado['total_receitas'] == Decimal('1750')
        assert resultado['total_despesas'] == Decimal('300')
        assert resultado['saldo_final'] == saldo + Decimal('1450')
        assert len(resultado['atrasados']) == 1
        assert resultado['projecao'][1]['receitas'] == Decimal('1750')


def test_tela_e_excel_iguais_e_dre_avisa_limitacao(caso_jumirim_ficticio):
    app, ids = caso_jumirim_ficticio
    with app.app_context():
        lanc(ids, 'Entrada futura', 'conta_receber', '425', date.today() + timedelta(days=1))
        lanc(ids, 'Saída futura', 'conta_pagar', '125', date.today() + timedelta(days=1))
        db.session.commit()
    client = app.test_client()
    client.post('/auth/login', data={'identificador':'admin_pr46','senha':'SenhaPr46Teste!'})
    capturados = []
    def capturar(sender, template, context, **extra):
        if template.name == 'financeiro/fluxo_caixa/dashboard.html':
            capturados.append(context)
    template_rendered.connect(capturar, app, weak=False)
    try:
        pagina = client.get(f'/financeiro/fluxo-caixa?conta_id={ids["conta"]}&periodo=30')
        assert pagina.status_code == 200 and len(capturados) == 1
        resposta = client.get(f'/financeiro/fluxo-caixa/exportar-excel?conta_id={ids["conta"]}&periodo=30')
        assert resposta.status_code == 200
        ws = openpyxl.load_workbook(BytesIO(resposta.data), data_only=True).active
        # Hoje é linha 6; amanhã é linha 7.
        assert Decimal(str(ws['B7'].value)) == capturados[0]['projecao'][1]['receitas']
        assert Decimal(str(ws['C7'].value)) == capturados[0]['projecao'][1]['despesas']
        assert client.get('/financeiro/dre?ano=2026&mes=9').status_code == 200
        assert 'DRE em revisão'.encode() in client.get('/financeiro/dre?ano=2026&mes=9').data
        xls = client.get('/financeiro/dre/exportar-excel?ano=2026&mes=9')
        assert 'EM REVISÃO' in openpyxl.load_workbook(BytesIO(xls.data)).active['A3'].value
    finally:
        template_rendered.disconnect(capturar, app)
